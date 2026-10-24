"""Read the official participants xlsx and export it for the Google Sheet.

    uv run tools/participants.py            # writes participants.csv (git-ignored)

Then in the Sheet: participants tab > File > Import > Upload > Replace current sheet.
"""

import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Participants-officiels-MT-2.xlsx"


def load_participants(path: Path = XLSX) -> list[dict]:
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(min_row=2, values_only=True)
    out = []
    for name, bib, *_ in rows:
        if bib is None or not name:
            continue
        out.append({"bib": int(bib), "name": str(name).strip(), "sex": ""})
    return out


def main():
    people = load_participants()
    bibs = [p["bib"] for p in people]
    dupes = {b for b in bibs if bibs.count(b) > 1}
    if dupes:
        print(f"WARNING duplicate bibs: {sorted(dupes)}")
    out = ROOT / "participants.csv"
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["bib", "name", "sex"])
        w.writeheader()
        w.writerows(people)
    print(f"{len(people)} participants -> {out} (bibs {min(bibs)}-{max(bibs)})")


if __name__ == "__main__":
    main()
